from pathlib import Path

import pytesseract
from pdf2image import convert_from_path
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from openpyxl import load_workbook
import pandas as pd


# ---------------------------------------------------------
# TESSERACT CONFIGURATION
# ---------------------------------------------------------

TESSERACT_PATH = Path(
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

if TESSERACT_PATH.exists():
    pytesseract.pytesseract.tesseract_cmd = str(TESSERACT_PATH)


# ---------------------------------------------------------
# SUPPORTED FILE TYPES
# ---------------------------------------------------------

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".pptx",
    ".xlsx",
    ".csv",
}


# ---------------------------------------------------------
# TXT
# ---------------------------------------------------------

def extract_txt(file_path: Path) -> str:
    return file_path.read_text(
        encoding="utf-8",
        errors="ignore"
    )


# ---------------------------------------------------------
# PDF - NORMAL + OCR
# ---------------------------------------------------------

def extract_pdf(file_path: Path) -> str:
    """
    First tries normal PDF text extraction.

    If no useful text is found, automatically converts
    PDF pages to images and performs OCR using Tesseract.
    """

    # ---------------------------------------------
    # STEP 1: NORMAL PDF TEXT EXTRACTION
    # ---------------------------------------------

    try:
        reader = PdfReader(str(file_path))

        text_parts = []

        for page in reader.pages:
            try:
                text = page.extract_text()

                if text and text.strip():
                    text_parts.append(text.strip())

            except Exception:
                continue

        normal_text = "\n".join(text_parts).strip()

        if normal_text:
            return normal_text

    except Exception:
        pass

    # ---------------------------------------------
    # STEP 2: OCR FOR SCANNED PDF
    # ---------------------------------------------

    try:

        images = convert_from_path(
            str(file_path),
            dpi=200
        )

        ocr_parts = []

        for page_number, image in enumerate(images, start=1):

            try:

                text = pytesseract.image_to_string(
                    image,
                    config="--psm 3"
                )

                if text and text.strip():
                    ocr_parts.append(
                        f"[Page {page_number}]\n{text.strip()}"
                    )

            except Exception:
                continue

        ocr_text = "\n\n".join(ocr_parts).strip()

        return ocr_text

    except Exception as error:
        raise RuntimeError(
            f"PDF OCR failed: {error}"
        )


# ---------------------------------------------------------
# DOCX
# ---------------------------------------------------------

def extract_docx(file_path: Path) -> str:

    document = DocxDocument(str(file_path))

    text_parts = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    row_values.append(text)

            if row_values:
                text_parts.append(
                    " | ".join(row_values)
                )

    return "\n".join(text_parts)


# ---------------------------------------------------------
# PPTX
# ---------------------------------------------------------

def extract_pptx(file_path: Path) -> str:

    presentation = Presentation(
        str(file_path)
    )

    text_parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_parts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

        if slide_parts:

            text_parts.append(
                f"[Slide {slide_number}]\n"
                + "\n".join(slide_parts)
            )

    return "\n\n".join(text_parts)


# ---------------------------------------------------------
# XLSX
# ---------------------------------------------------------

def extract_xlsx(file_path: Path) -> str:

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    text_parts = []

    for worksheet in workbook.worksheets:

        sheet_parts = []

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is not None:

                    values.append(
                        str(value).strip()
                    )

            if values:

                sheet_parts.append(
                    " | ".join(values)
                )

        if sheet_parts:

            text_parts.append(
                f"[Sheet: {worksheet.title}]\n"
                + "\n".join(sheet_parts)
            )

    workbook.close()

    return "\n\n".join(text_parts)


# ---------------------------------------------------------
# CSV
# ---------------------------------------------------------

def extract_csv(file_path: Path) -> str:

    dataframe = pd.read_csv(
        file_path,
        dtype=str,
        keep_default_na=False
    )

    return dataframe.to_csv(
        index=False
    )


# ---------------------------------------------------------
# MAIN TEXT EXTRACTION
# ---------------------------------------------------------

def extract_text(file_path: Path) -> str:

    extension = file_path.suffix.lower()

    if extension == ".txt":
        return extract_txt(file_path)

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    if extension == ".xlsx":
        return extract_xlsx(file_path)

    if extension == ".csv":
        return extract_csv(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


# ---------------------------------------------------------
# WORD COUNT
# ---------------------------------------------------------

def count_words(text: str) -> int:

    if not text:
        return 0

    return len(text.split())


# ---------------------------------------------------------
# FILE METADATA
# ---------------------------------------------------------

def get_file_metadata(
    file_path: Path,
    text: str
) -> dict:

    return {
        "filename": file_path.name,
        "extension": file_path.suffix.lower(),
        "size": file_path.stat().st_size,
        "word_count": count_words(text),
        "character_count": len(text),
    }